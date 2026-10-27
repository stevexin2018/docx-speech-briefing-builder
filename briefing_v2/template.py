"""Word styles and minimal table enhancements, separate from content conversion."""
import re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .word_compatibility import prepare_word_layout


def configure_reference(path):
    document = Document(path)
    for section in document.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(2)
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(10.5)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
        style.paragraph_format.line_spacing = 1.2
        style.paragraph_format.space_after = Pt(5)
    for name in ["Title", *[f"Heading {level}" for level in range(1, 10)]]:
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(18 if name == "Title" else {"Heading 1": 16, "Heading 2": 13}.get(name, 11))
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "黑体")
        style.paragraph_format.keep_with_next = True
    document.save(path)


def enhance_document(path, title=None, topic_id=None, default_template=True, add_title=False):
    document = Document(path)
    if title:
        document.core_properties.title = title
    if topic_id:
        document.core_properties.identifier = topic_id
    if add_title:
        paragraph = document.add_paragraph(title, style="Title")
        document.element.body.remove(paragraph._p)
        document.element.body.insert(0, paragraph._p)
    layout = prepare_word_layout(document)
    for table in document.tables:
        if not table.rows:
            continue
        properties = table.rows[0]._tr.get_or_add_trPr()
        if properties.find(qn("w:tblHeader")) is None:
            properties.append(OxmlElement("w:tblHeader"))
        for row in table.rows:
            tr_properties = row._tr.get_or_add_trPr()
            if tr_properties.find(qn("w:cantSplit")) is None:
                tr_properties.append(OxmlElement("w:cantSplit"))
        # Repeated headings must travel with a data row, not stand alone at the
        # foot of the preceding page. Do not chain the whole multi-page table.
        for cell in table.rows[0].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
                paragraph.paragraph_format.keep_together = True
        if default_template:
            borders = OxmlElement("w:tblBorders")
            for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
                edge = OxmlElement(f"w:{side}")
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    edge.set(qn(f"w:{key}"), value)
                borders.append(edge)
            table._tbl.tblPr.append(borders)
    # Apply pagination to the output, including custom references. A fix to one
    # trial's template must not be required by every subsequent report.
    def formula(element):
        if element.xpath('.//m:oMath'):
            return True
        if element.tag != qn('w:p'):
            return False
        value = ''.join(element.xpath('.//w:t/text()')).strip()
        return (len(value) < 150 and re.match(r'^[A-Za-zα-ωΑ-Ω][\w]*\s*[=≈]', value)
                and re.search(r'[·⋅×/^()]', value))

    body = list(document.element.body)
    for index, element in enumerate(body):
        if element.tag != qn("w:p"):
            continue
        following = body[index + 1] if index + 1 < len(body) else None
        is_math = formula(element)
        next_math = following is not None and formula(following)
        if is_math or next_math:
            from docx.text.paragraph import Paragraph
            paragraph = Paragraph(element, document._body)
            paragraph.paragraph_format.keep_together = True
            if next_math:
                paragraph.paragraph_format.keep_with_next = True
    document.save(path)
    return layout
