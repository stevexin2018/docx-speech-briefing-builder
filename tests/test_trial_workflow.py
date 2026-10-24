"""Cache isolation, recorded failures and preserved output contracts."""
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest
from unittest.mock import patch

from docx import Document

import briefing_workflow as workflow
from build_briefing import build_docx_worker
from briefing_v2.engine import find_pandoc

try:
    PANDOC = find_pandoc()
except FileNotFoundError:
    PANDOC = None


class TrialFixture:
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.root = Path(self.folder.name)
        self.source = self.root / '源稿.md'
        self.source.write_text('# Report\n\n正文与 `C:\\Temp\\_data`。', encoding='utf-8')
        self.word, self.audio = self.root / '报告.docx', self.root / '报告.mp3'

    def tearDown(self):
        self.folder.cleanup()

    def build(self, **kwargs):
        return workflow.build_cached_briefing(self.source, self.word, engine='legacy', **kwargs)


class TrialTests(TrialFixture, unittest.TestCase):
    def test_unsupported_speech_math_preserves_audio_and_records_failure_before_tts(self):
        self.source.write_text(r'$$x=\notARealCommand{a}$$', encoding='utf-8')
        self.audio.write_bytes(b'previous-audio')
        with patch('clean_speech_text.generate_speech_audio') as synthesize:
            with self.assertRaises(workflow.CachedBuildError) as caught:
                workflow.build_cached_briefing(self.source, mp3=self.audio)
        synthesize.assert_not_called()
        self.assertEqual(self.audio.read_bytes(), b'previous-audio')
        report = caught.exception.report
        self.assertEqual(report['acceptance']['technical_build'], 'failed')
        self.assertIn('UNSUPPORTED_SPEECH_MATH', [d['code'] for d in report['diagnostics']])

    def test_success_does_not_claim_visual_listening_or_engineering_acceptance(self):
        report = self.build()
        self.assertEqual(report['status'], 'success')
        self.assertTrue(report['build_contract_passed'])
        self.assertTrue(report['review_required'])
        self.assertTrue(report['manual_review_required'])
        self.assertFalse(report['fully_accepted'])
        self.assertEqual(report['acceptance']['visual'], 'not_performed')
        self.assertEqual(report['acceptance']['listening'], 'not_requested')
        self.assertEqual(report['acceptance']['engineering'], 'not_assessed')

    def test_cache_hit_has_a_new_record_and_restored_diagnostics(self):
        first = self.build()
        timestamp = self.word.stat().st_mtime_ns
        Path(first['word_diagnostics_path']).unlink()
        second = self.build()
        self.assertEqual(second['reused'], ['docx'])
        self.assertNotEqual(first['run_id'], second['run_id'])
        self.assertTrue(Path(second['word_diagnostics_path']).exists())
        self.assertEqual(self.word.stat().st_mtime_ns, timestamp)
        self.assertEqual(json.loads(Path(first['record_path']).read_text(encoding='utf-8'))['status'], 'success')

    def test_corrupt_output_is_rebuilt(self):
        self.build()
        self.word.write_bytes(b'broken')
        result = self.build()
        self.assertEqual(result['rebuilt'], ['docx'])
        self.assertTrue(workflow.valid_asset(self.word))

    def test_invalid_cache_metadata_is_recorded_and_rebuilt(self):
        self.word.with_suffix('.assets.json').write_text('{broken')
        result = self.build()
        self.assertEqual(result['rebuilt'], ['docx'])
        self.assertIn('CACHE_METADATA_INVALID', [d['code'] for d in result['diagnostics']])

    def test_source_change_blocks_delivery(self):
        def worker(source, output, title, topic):
            self.assertEqual(Path(source).read_text(encoding='utf-8'), self.source.read_text(encoding='utf-8'))
            self.source.write_text('modified', encoding='utf-8')
            document = Document()
            document.add_paragraph('frozen content')
            document.save(output)
        with self.assertRaises(workflow.CachedBuildError) as caught:
            self.build(workers={'docx': worker})
        self.assertFalse(caught.exception.report['ready_for_delivery'])
        self.assertIn('SOURCE_CHANGED', [d['code'] for d in caught.exception.report['diagnostics']])

    def test_busy_output_preserves_manifest_and_records_the_attempt(self):
        first = self.build()
        manifest = Path(first['manifest_path'])
        previous = manifest.read_bytes()
        with workflow._lock(self.word.with_name('.' + self.word.name + '.briefing.lock')):
            with self.assertRaises(workflow.CachedBuildError) as caught:
                self.build()
        report = caught.exception.report
        self.assertEqual(manifest.read_bytes(), previous)
        self.assertEqual(report['diagnostics'][0]['code'], 'OUTPUT_BUSY')
        self.assertTrue(Path(report['record_path']).is_file())
        self.assertTrue(Path(report['word_diagnostics_path']).is_file())

    def test_reserved_lock_cannot_be_used_as_source(self):
        source = self.word.with_name('.' + self.word.name + '.briefing.lock')
        source.write_text('preserve')
        with self.assertRaises(ValueError):
            workflow.build_cached_briefing(source, self.word, engine='legacy')
        self.assertEqual(source.read_text(), 'preserve')

    def test_word_only_and_audio_only_share_the_manifest_lock(self):
        first = self.build()
        manifest = Path(first['manifest_path'])
        previous = manifest.read_bytes()
        with workflow._lock(manifest.with_name('.' + manifest.name + '.briefing.lock')):
            with self.assertRaises(workflow.CachedBuildError) as caught:
                workflow.build_cached_briefing(self.source, mp3=self.audio)
        self.assertEqual(caught.exception.report['diagnostics'][0]['code'], 'OUTPUT_BUSY')
        self.assertEqual(manifest.read_bytes(), previous)
        self.assertFalse(self.audio.exists())

    @unittest.skipUnless(PANDOC, 'Pandoc required')
    def test_missing_pandoc_failure_has_diagnostics_and_retains_previous_word(self):
        self.build()
        previous = self.word.read_bytes()
        with self.assertRaises(workflow.CachedBuildError) as caught:
            workflow.build_cached_briefing(self.source, self.word, pandoc_path=str(self.root / 'missing-pandoc'))
        report = caught.exception.report
        self.assertEqual(self.word.read_bytes(), previous)
        self.assertIn('DEPENDENCY_NOT_FOUND', [d['code'] for d in report['diagnostics']])
        self.assertTrue(Path(report['word_diagnostics_path']).is_file())

    @unittest.skipUnless(PANDOC, 'Pandoc required')
    def test_unknown_formula_is_located_and_failed_attempt_is_saved(self):
        self.build()
        previous = self.word.read_bytes()
        self.source.write_text('# Report\n\n$\\unknowntrialcommand{x}$', encoding='utf-8')
        with self.assertRaises(workflow.CachedBuildError) as caught:
            workflow.build_cached_briefing(self.source, self.word)
        report = caught.exception.report
        self.assertEqual(self.word.read_bytes(), previous)
        diagnostic = next(d for d in report['diagnostics'] if d['code'] == 'MATH_CONVERSION_FAILED')
        self.assertEqual(diagnostic['source'], str(self.source))
        self.assertEqual(diagnostic['line'], 3)
        self.assertEqual(json.loads(Path(report['record_path']).read_text(encoding='utf-8'))['status'], 'failed')

    @unittest.skipUnless(PANDOC, 'Pandoc required')
    def test_cached_warnings_are_rebound_to_current_source_path(self):
        self.source.write_text('# Report\n\n5^', encoding='utf-8')
        first = workflow.build_cached_briefing(self.source, self.word, title='fixed title')
        other = self.root / 'other.md'
        other.write_bytes(self.source.read_bytes())
        second = workflow.build_cached_briefing(other, self.word, title='fixed title')
        self.assertTrue(first['review_required'])
        self.assertEqual(second['reused'], ['docx'])
        self.assertTrue(second['review_required'])
        self.assertTrue(all(d['source'] == str(other) for d in second['diagnostics']))

    @unittest.skipUnless(PANDOC, 'Pandoc required')
    def test_changed_reference_rebuilds_word(self):
        reference = self.root / 'template.docx'
        Document().save(reference)
        workflow.build_cached_briefing(self.source, self.word, reference_doc=reference)
        document = Document(reference)
        document.styles['Normal'].font.name = 'Arial'
        document.save(reference)
        result = workflow.build_cached_briefing(self.source, self.word, reference_doc=reference)
        self.assertEqual(result['rebuilt'], ['docx'])

    def test_corrupt_or_formula_only_word_is_validated_structurally(self):
        from docx.oxml import parse_xml
        document = Document()
        document.add_paragraph()._p.append(parse_xml(
            '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
            '<m:r><m:t>x</m:t></m:r></m:oMath>'))
        document.save(self.word)
        self.assertTrue(workflow.valid_asset(self.word))


@unittest.skipUnless(shutil.which('ffmpeg') and shutil.which('ffprobe'), 'FFmpeg tools required')
class AudioCacheTests(TrialFixture, unittest.TestCase):
    def setUp(self):
        super().setUp()
        self.fixture = self.root / 'fixture.mp3'
        subprocess.run(['ffmpeg', '-v', 'error', '-y', '-f', 'lavfi', '-i',
                        'sine=duration=0.1', str(self.fixture)], check=True, capture_output=True)
        self.calls = {'docx': 0, 'mp3': 0}

    def word_worker(self, *args, **kwargs):
        self.calls['docx'] += 1
        return build_docx_worker(*args, **kwargs)

    def audio_worker(self, source, output, rate, voice, proxy):
        self.calls['mp3'] += 1
        shutil.copy2(self.fixture, output)

    def cached(self, **kwargs):
        return workflow.build_cached_briefing(self.source, self.word, self.audio,
            workers={'docx': self.word_worker, 'mp3': self.audio_worker}, **kwargs)

    @unittest.skipUnless(PANDOC, 'Pandoc required')
    def test_engine_and_template_changes_reuse_audio(self):
        self.cached(engine='legacy')
        second = self.cached(engine='pandoc')
        self.assertEqual(second['rebuilt'], ['docx'])
        self.assertEqual(second['reused'], ['mp3'])
        template = self.root / 'template.docx'
        Document().save(template)
        self.cached(reference_doc=template)
        self.assertEqual(self.calls, {'docx': 3, 'mp3': 1})

    def test_rate_and_voice_changes_only_rebuild_audio(self):
        self.cached(engine='legacy')
        second = self.cached(engine='legacy', rate='+10%')
        third = self.cached(engine='legacy', rate='+10%', voice='zh-CN-XiaoyiNeural')
        self.assertEqual(second['rebuilt'], ['mp3'])
        self.assertEqual(third['rebuilt'], ['mp3'])
        self.assertEqual(self.calls, {'docx': 1, 'mp3': 3})

    def test_core_release_change_rebuilds_word_and_reuses_unchanged_audio(self):
        self.cached(engine='legacy')
        with patch.object(workflow, '__version__', 'next-preview'):
            result = self.cached(engine='legacy')
        self.assertEqual(result['rebuilt'], ['docx'])
        self.assertEqual(result['reused'], ['mp3'])

    def test_audio_failure_is_soft_only_when_requested_and_resume_reuses_word(self):
        def fail(*args):
            raise RuntimeError('https://user:secret@proxy.invalid/token')
        first = workflow.build_cached_briefing(self.source, self.word, self.audio, engine='legacy',
            workers={'docx': self.word_worker, 'mp3': fail}, audio_optional=True)
        self.assertEqual(first['delivery_mode'], 'word_only')
        self.assertIn('docx', first['assets'])
        self.assertNotIn('mp3', first['assets'])
        self.assertNotIn('secret', Path(first['record_path']).read_text(encoding='utf-8'))
        second = self.cached(engine='legacy', audio_optional=True)
        self.assertEqual(second['rebuilt'], ['mp3'])
        self.assertNotIn('delivery_mode', second)
        self.assertEqual(self.calls['docx'], 1)

    def test_corrupt_audio_rebuilds_without_word(self):
        self.cached(engine='legacy')
        self.audio.write_bytes(b'invalid mp3')
        second = self.cached(engine='legacy')
        self.assertEqual(second['rebuilt'], ['mp3'])

    def test_audio_only_never_probes_pandoc(self):
        with patch('briefing_workflow._word_config', side_effect=AssertionError('Word was called')):
            result = workflow.build_cached_briefing(self.source, mp3=self.audio,
                workers={'mp3': self.audio_worker})
        self.assertIsNone(result['engine'])
        self.assertIsNone(result['word_diagnostics_path'])
        self.assertEqual(result['rebuilt'], ['mp3'])

    def test_tts_authorization_failure_keeps_safe_actionable_diagnostic(self):
        def fail(*args):
            raise RuntimeError('TTS authorization rejected: HTTP 403 https://user:secret@host')
        with self.assertRaises(workflow.CachedBuildError) as caught:
            workflow.build_cached_briefing(self.source, mp3=self.audio, workers={'mp3': fail})
        diagnostic = caught.exception.report['diagnostics'][0]
        self.assertEqual(diagnostic['code'], 'TTS_AUTH_REJECTED')
        self.assertIn('403', diagnostic['message'])
        self.assertNotIn('secret', Path(caught.exception.report['record_path']).read_text(encoding='utf-8'))
