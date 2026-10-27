"""Compatibility and layout checks; these do not emulate mobile Word."""
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from lxml import etree

from briefing_v2.engine import find_pandoc, render_markdown
from briefing_v2.word_compatibility import prepare_word_layout, WORD_URI

M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
FRACTION = f'<m:oMath xmlns:m="{M}"><m:f><m:num><m:r><m:t>a</m:t></m:r></m:num><m:den><m:r><m:t>b</m:t></m:r></m:den></m:f></m:oMath>'


def formula_xml(document):
    return [etree.tostring(element, method='c14n', exclusive=True)
            for element in document.element.xpath('.//m:oMath')]


class WordCompatibilityTests(unittest.TestCase):
    def test_modern_mode_replaces_old_mode_without_losing_other_settings(self):
        document = Document()
        compatibility = document.settings.element.find(qn('w:compat'))
        if compatibility is None:
            compatibility = OxmlElement('w:compat')
            document.settings.element.append(compatibility)
        extra = OxmlElement('w:compatSetting')
        for name, value in [('name', 'retainCustomSetting'), ('uri', WORD_URI), ('val', '1')]:
            extra.set(qn('w:' + name), value)
        compatibility.append(extra)
        for value in ('12', '14'):
            item = OxmlElement('w:compatSetting')
            item.set(qn('w:name'), 'compatibilityMode')
            item.set(qn('w:uri'), WORD_URI)
            item.set(qn('w:val'), value)
            compatibility.append(item)
        prepare_word_layout(document)
        modes = [c for c in compatibility if c.get(qn('w:name')) == 'compatibilityMode']
        self.assertEqual(len(modes), 1)
        self.assertEqual(modes[0].get(qn('w:val')), '15')
        self.assertIn(extra, list(compatibility))

    def test_existing_math_font_and_wrap_indent_are_retained(self):
        document = Document()
        for existing in document.settings.element.findall(qn('m:mathPr')):
            document.settings.element.remove(existing)
        properties = OxmlElement('m:mathPr')
        for name, value in [('mathFont', 'STIX Two Math'), ('wrapIndent', '720')]:
            item = OxmlElement('m:' + name)
            item.set(qn('m:val'), value)
            properties.append(item)
        document.settings.element.append(properties)
        prepare_word_layout(document)
        self.assertEqual(properties.find(qn('m:mathFont')).get(qn('m:val')), 'STIX Two Math')
        self.assertEqual(properties.find(qn('m:wrapIndent')).get(qn('m:val')), '720')
        self.assertEqual(properties.find(qn('m:brkBin')).get(qn('m:val')), 'before')

    def test_standalone_complex_inline_equation_and_caption_are_separated(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph._p.append(parse_xml(FRACTION))
        paragraph.add_run(' 【待校验公式模型】').bold = True
        before = formula_xml(document)
        changes = prepare_word_layout(document)
        self.assertEqual(formula_xml(document), before)
        self.assertEqual(changes['inline_equations_promoted'], 1)
        self.assertEqual(changes['mixed_paragraphs_separated'], 1)
        self.assertEqual(len(document.paragraphs), 2)
        self.assertTrue(document.paragraphs[0]._p.xpath('./m:oMathPara'))
        self.assertEqual(document.paragraphs[1].text.strip(), '【待校验公式模型】')
        self.assertTrue(document.paragraphs[1].runs[0].bold)
        self.assertTrue(document.paragraphs[0].paragraph_format.keep_with_next)
        self.assertEqual(document.paragraphs[0]._p.xpath('./w:pPr/w:spacing/@w:lineRule'), ['auto'])

    def test_inline_equations_in_prose_and_table_cells_are_unchanged(self):
        document = Document()
        paragraph = document.add_paragraph('压力 ')
        paragraph._p.append(parse_xml(FRACTION))
        paragraph.add_run(' 用于说明。')
        cell = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
        cell._p.append(parse_xml(FRACTION))
        before = formula_xml(document)
        prepare_word_layout(document)
        self.assertEqual(formula_xml(document), before)
        self.assertTrue(paragraph._p.xpath('./m:oMath'))
        self.assertTrue(cell._p.xpath('./m:oMath'))

    def test_numbering_bookmarks_and_section_breaks_are_not_relocated(self):
        document = Document()
        for special in ('w:numPr', 'w:sectPr'):
            paragraph = document.add_paragraph()
            paragraph._p.get_or_add_pPr().append(OxmlElement(special))
            paragraph._p.append(parse_xml(FRACTION))
        bookmark = document.add_paragraph()
        node = OxmlElement('w:bookmarkStart')
        node.set(qn('w:id'), '42')
        node.set(qn('w:name'), 'equation')
        bookmark._p.append(node)
        bookmark._p.append(parse_xml(FRACTION))
        before = [etree.tostring(p._p) for p in document.paragraphs]
        prepare_word_layout(document)
        self.assertEqual([etree.tostring(p._p) for p in document.paragraphs], before)

    def test_display_math_mixed_with_text_preserves_all_text_and_math(self):
        document = Document()
        paragraph = document.add_paragraph('公式说明')
        display = OxmlElement('m:oMathPara')
        display.append(parse_xml(FRACTION))
        paragraph._p.append(display)
        paragraph.add_run('【备注】')
        before = formula_xml(document)
        prepare_word_layout(document)
        self.assertEqual(formula_xml(document), before)
        self.assertEqual([p.text for p in document.paragraphs], ['公式说明', '', '【备注】'])

    def test_reapplying_layout_is_idempotent(self):
        document = Document()
        document.add_paragraph()._p.append(parse_xml(FRACTION))
        prepare_word_layout(document)
        before = etree.tostring(document.element), etree.tostring(document.settings.element)
        prepare_word_layout(document)
        self.assertEqual((etree.tostring(document.element), etree.tostring(document.settings.element)), before)

    def test_trailing_status_moves_outside_math_without_changing_algebra(self):
        document = Document()
        paragraph = document.add_paragraph()
        display = OxmlElement('m:oMathPara')
        equation = parse_xml(FRACTION)
        before = etree.tostring(equation, method='c14n', exclusive=True)
        for value in ('\u2001', '【待核验公式模型】'):
            run = OxmlElement('m:r')
            text = OxmlElement('m:t')
            text.text = value
            run.append(text)
            equation.append(run)
        display.append(equation)
        paragraph._p.append(display)
        changes = prepare_word_layout(document)
        self.assertEqual(formula_xml(document), [before])
        self.assertEqual(changes['status_annotations_moved'], 1)
        self.assertEqual(document.paragraphs[1].text, '【待核验公式模型】')

    def test_mathematical_labels_placeholders_and_mixed_runs_are_not_extracted(self):
        for tail in ('[x]', '1.0 【待核验】'):
            document = Document()
            equation = parse_xml(FRACTION)
            run = OxmlElement('m:r')
            value = OxmlElement('m:t')
            value.text = tail
            run.append(value)
            equation.append(run)
            display = OxmlElement('m:oMathPara')
            display.append(equation)
            document.add_paragraph()._p.append(display)
            before = formula_xml(document)
            self.assertEqual(prepare_word_layout(document)['status_annotations_moved'], 0)
            self.assertEqual(formula_xml(document), before)
        document = Document()
        equation = parse_xml(f'<m:oMath xmlns:m="{M}"><m:r><m:t>x=</m:t></m:r><m:r><m:t>【待核验】</m:t></m:r></m:oMath>')
        display = OxmlElement('m:oMathPara')
        display.append(equation)
        document.add_paragraph()._p.append(display)
        before = formula_xml(document)
        self.assertEqual(prepare_word_layout(document)['status_annotations_moved'], 0)
        self.assertEqual(formula_xml(document), before)


try:
    PANDOC = find_pandoc()
except FileNotFoundError:
    PANDOC = None


@unittest.skipUnless(PANDOC, 'Pandoc is required')
class NativeOutputCompatibilityTests(unittest.TestCase):
    def test_generated_docx_has_modern_settings_and_native_formula_integrity(self):
        with tempfile.TemporaryDirectory() as folder:
            output, reference = Path(folder) / 'report.docx', Path(folder) / 'reference.docx'
            Document().save(reference)
            reference_before = reference.read_bytes()
            source = '# Review\n\n$L=\\frac{a}{b}\\sqrt{x}$ **【待校验】**\n\n路径 `C:\\Temp\\_data`。'
            report = render_markdown(source, output, pandoc_path=PANDOC, reference_doc=reference)
            document = Document(output)
            self.assertEqual(report['output_counts']['math'], 1)
            self.assertEqual(report['word_layout']['compatibility_mode'], 15)
            self.assertTrue(document.element.xpath('.//m:f'))
            self.assertTrue(document.element.xpath('.//m:rad'))
            modes = document.settings.element.xpath('./w:compat/w:compatSetting[@w:name="compatibilityMode"]')
            self.assertEqual([mode.get(qn('w:val')) for mode in modes], ['15'])
            self.assertEqual(reference.read_bytes(), reference_before)
            self.assertTrue(any('C:\\Temp\\_data' in p.text for p in document.paragraphs))
