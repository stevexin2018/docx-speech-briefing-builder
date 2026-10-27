"""The verified file-format fix must not restructure mathematical content."""
from copy import deepcopy
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from lxml import etree

from briefing_v2.compatibility import use_modern_word_mode, WORD_URI
from briefing_v2.engine import find_pandoc, render_markdown


class CompatibilityModeTests(unittest.TestCase):
    def test_only_compatibility_mode_changes_and_is_idempotent(self):
        document = Document()
        paragraph = document.add_paragraph('Keep this formula and label: ')
        paragraph._p.append(parse_xml('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:f><m:num><m:r><m:t>a</m:t></m:r></m:num><m:den><m:r><m:t>b</m:t></m:r></m:den></m:f><m:r><m:t>【待核验】</m:t></m:r></m:oMath>'))
        body_before = etree.tostring(document.element)
        compatibility = document.settings.element.find(qn('w:compat'))
        for existing in compatibility.findall(qn('w:compatSetting')):
            if existing.get(qn('w:name')) == 'compatibilityMode':
                compatibility.remove(existing)
        old_mode = OxmlElement('w:compatSetting')
        for name, value in [('name', 'compatibilityMode'), ('uri', WORD_URI), ('val', '12')]:
            old_mode.set(qn('w:' + name), value)
        compatibility.append(old_mode)
        expected = deepcopy(document.settings.element)
        expected.xpath('./w:compat/w:compatSetting[@w:name="compatibilityMode"]')[0].set(qn('w:val'), '15')
        self.assertEqual(use_modern_word_mode(document), {'compatibility_mode': 15})
        self.assertEqual(etree.tostring(document.element), body_before)
        self.assertEqual(etree.tostring(document.settings.element), etree.tostring(expected))
        after = etree.tostring(document.settings.element)
        use_modern_word_mode(document)
        self.assertEqual(etree.tostring(document.settings.element), after)

    def test_missing_compat_and_duplicate_old_modes_are_handled(self):
        for missing in (True, False):
            document = Document()
            settings = document.settings.element
            compat = settings.find(qn('w:compat'))
            if missing:
                settings.remove(compat)
            else:
                for value in ('12', '14'):
                    mode = OxmlElement('w:compatSetting')
                    mode.set(qn('w:name'), 'compatibilityMode')
                    mode.set(qn('w:val'), value)
                    compat.append(mode)
                other = OxmlElement('w:compatSetting')
                other.set(qn('w:name'), 'retainCustomSetting')
                other.set(qn('w:val'), '1')
                compat.append(other)
            use_modern_word_mode(document)
            modes = settings.xpath('./w:compat/w:compatSetting[@w:name="compatibilityMode"]')
            self.assertEqual(len(modes), 1)
            self.assertEqual(modes[0].get(qn('w:val')), '15')
            if not missing:
                self.assertIn(other, list(compat))
            else:
                self.assertLess(list(settings).index(settings.find(qn('w:compat'))),
                                list(settings).index(settings.find(qn('m:mathPr'))))


try:
    PANDOC = find_pandoc()
except FileNotFoundError:
    PANDOC = None


@unittest.skipUnless(PANDOC, 'Pandoc is required')
class NativeBuildCompatibilityTests(unittest.TestCase):
    def test_native_build_records_mode_and_preserves_reference_and_formula_label(self):
        with tempfile.TemporaryDirectory() as folder:
            reference, output = Path(folder) / 'reference.docx', Path(folder) / 'output.docx'
            Document().save(reference)
            before = reference.read_bytes()
            source = '# Review\n\n$$L=\\frac{a}{b}\\sqrt{x}\\quad\\text{【待核验】}$$\n'
            for custom_reference in (None, reference):
                with self.subTest(custom_reference=bool(custom_reference)):
                    report = render_markdown(source, output, pandoc_path=PANDOC, reference_doc=custom_reference)
                    self.assertEqual(report['word_layout'], {'compatibility_mode': 15})
                    self.assertEqual(report['output_counts']['math'], 1)
                    document = Document(output)
                    self.assertTrue(document.element.xpath('.//m:f'))
                    self.assertTrue(document.element.xpath('.//m:rad'))
                    self.assertIn('【待核验】', ''.join(document.element.xpath('.//m:t/text()')))
            self.assertEqual(reference.read_bytes(), before)
