"""Public, minimal reproductions derived from the three private trial reports."""
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from clean_speech_text import clean_speech_text
from engineering_text import normalize_quantities
from speech_structure import SpeechMathError, speak_math
from briefing_v2.engine import find_pandoc, render_markdown
from docx import Document
from docx.oxml.ns import qn

try:
    PANDOC = find_pandoc()
except FileNotFoundError:
    PANDOC = None


class SpeechStructureTests(unittest.TestCase):
    def test_angles_temperatures_and_star_superscripts_in_context(self):
        cases = {
            r'$45^\circ$': ('四十五 度',),
            r'$1.5 \times 45^\circ$': ('一点五 乘以 四十五 度',),
            r'$80\ ^\circ\text{C}$': ('80 摄氏度',),
            r'$-29\ ^\circ\text{C} \text{ @ } 0.6\text{ MPa}$': ('零下 二十九 摄氏度', '在 零点六 兆帕'),
            r'$T=20^{\circ}\mathrm{C}$': ('T 等于 二十 摄氏度',),
            r'$S^{\ast}$': ('S 上标星号',),
            r'$f\circ g$': ('f 复合 g',),
            r'$x^2$': ('x 的平方',),
        }
        for source, expected in cases.items():
            with self.subTest(source=source):
                spoken = clean_speech_text(source)
                for value in expected:
                    self.assertIn(value, spoken)
                self.assertNotIn('次方 度', spoken)
        with self.assertRaises(SpeechMathError):
            speak_math(r'x^\circ')

    def test_fraction_retains_the_complete_denominator_and_minus(self):
        value = clean_speech_text(r'$$t=\frac{PR}{SE-0.6P}$$')
        self.assertIn('分子为 P 乘以 R', value)
        self.assertIn('除以分母 S 乘以 E 减去 零点六 乘以 P，分式结束', value)

    def test_addition_power_and_subscript_survive_the_prose_cleaner(self):
        value = clean_speech_text(r'$$P_i=P_0+\frac{\rho g h_i}{10^6}$$')
        self.assertIn('加上', value)
        self.assertIn('十 的 六 次方', value)
        self.assertIn('下标 i，下标结束', value)
        self.assertNotIn('10 6', value)

    def test_nested_fraction_scripts_and_trigonometry(self):
        for source in [r'$$s_g=\frac{l_{\mathrm{slot}}-d_b}{2}$$',
                       r'$$\ell=s-\frac{d_1+d_2}{2}$$',
                       r'$$T=\frac{\varphi m g}{n\cos\beta}$$']:
            value = clean_speech_text(source)
            self.assertNotRegex(value, r'\\[A-Za-z]+|\b(?:frac|ell|varphi|cos)\b')
            self.assertIn('分式结束', value)
        self.assertIn('减去', clean_speech_text(r'$$\ell=s-\frac{d_1+d_2}{2}$$'))
        self.assertIn('除以分母 n 乘以 余弦 贝塔', clean_speech_text(r'$$T=\frac{\varphi m g}{n\cos\beta}$$'))

    def test_nested_roots_and_matrix(self):
        self.assertIn('根号结束', speak_math(r'\sqrt{\frac{a^2}{b^2}}'))
        matrix = speak_math(r'\begin{bmatrix}1&2\\3&4\end{bmatrix}')
        self.assertIn('第一行 一，二', matrix)
        self.assertIn('第二行 三，四', matrix)

    def test_unsupported_math_fails_explicitly(self):
        with self.assertRaises(SpeechMathError):
            clean_speech_text(r'$$x=\notARealCommand{a}$$')
        with self.assertRaises(SpeechMathError):
            clean_speech_text('$$x=1')

    def test_container_metadata_does_not_become_narration(self):
        value = clean_speech_text('::: {custom-style="FormulaIntro"}\n完整正文\n:::')
        self.assertIn('完整正文', value)
        self.assertNotIn('FormulaIntro', value)
        self.assertNotIn('custom', value)

    def test_hole_count_and_standard_identifiers(self):
        self.assertEqual(clean_speech_text('2-24×64'), '2 个 24 乘以 64。')
        standard = clean_speech_text('JB/T4712.1-2007')
        self.assertIn('年版 二零零七', standard)
        self.assertNotIn('除以', standard)
        self.assertNotIn('至', standard)
        self.assertIn('连接号 F', clean_speech_text('BⅡ2000-F'))
        self.assertNotIn('除以', clean_speech_text('BⅡ2000-F / BⅡ2000-S'))
        self.assertIn('或', clean_speech_text('BⅡ2000-F/S'))
        self.assertEqual(clean_speech_text('2-Φ24'), '2 个直径 24。')
        self.assertEqual(clean_speech_text('10-20mm'), '10至20 毫米。')

    def test_identifiers_are_not_litre_units(self):
        for source in ['∠50×50×6，L=2000，两件', '∠50×50×6，L = 2000，两件', '143237-07，A1/M1与L2～4', 'Φ610/620，t16，L115，Q345R']:
            self.assertEqual(normalize_quantities(source), source)
        self.assertEqual(normalize_quantities('10,mm3'), '10 mm³')
        self.assertEqual(normalize_quantities('10,L'), '10 L')

    def test_paths_and_code_are_not_interpreted_as_equations(self):
        value = clean_speech_text(r'`C:\项目\报告.docx`')
        self.assertIn('C 盘，项目，报告 点 docx', value)
        self.assertNotIn('下标', clean_speech_text('`x_i^2`'))
        self.assertNotIn('等于', clean_speech_text(r'C:\data\$x$\report.docx'))

    def test_labelled_section_sketch_and_identifier_speech(self):
        sketch = ('```\n[48 mm SA-105 法兰盖基体]\n'
                  '================================== <-- 结构承载受压面\n'
                  '[4 mm 硬橡胶衬里] <-- 不承压\n'
                  '----------------------------------\n'
                  '[3 mm 软垫片]\n```')
        spoken = clean_speech_text(sketch)
        self.assertIn('结构承载受压面', spoken)
        self.assertIn('不承压', spoken)
        self.assertNotIn('====', spoken)
        self.assertNotIn('----', spoken)
        self.assertNotIn('等于', spoken)
        self.assertIn('W 下标 m 二，下标结束', clean_speech_text(r'$W_{m2}$'))
        self.assertNotIn('乘以 二，下标结束', clean_speech_text(r'$W_{m2}$'))
        self.assertIn('直径 三千六百 毫米', clean_speech_text(r'$\Phi 3600\text{ mm}$'))
        self.assertIn('R z 五十 至 八十 微米', clean_speech_text(r'$Rz\ 50 \sim 80\ \mu\text{m}$'))
        self.assertIn('派 乘以 G 乘以 m 乘以 P', clean_speech_text(r'$\pi G m P$'))
        self.assertIn('P 毫米', clean_speech_text(r'$P\text{ mm}$'))
        self.assertIn('二点五 英寸', clean_speech_text(r'$\le 2.5"$'))
        self.assertIn('x = 1', clean_speech_text('```\nx = 1\n```'))


@unittest.skipUnless(PANDOC, 'Pandoc is required for native Word regression')
class NativeWordReleaseTests(unittest.TestCase):
    def test_section_sketch_removes_rules_but_keeps_labels_and_inches(self):
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / 'sketch.docx'
            source = ('# Review\n\n```\n[48 mm 法兰盖]\n'
                      '=================== <-- 承压面\n'
                      '[4 mm 衬里] <-- 不承压\n'
                      '-------------------\n[3 mm 垫片]\n```\n\n'
                      '螺柱直径为 $\\le 2.5"$ 或 $> 2.5"$。')
            report = render_markdown(source, output, pandoc_path=PANDOC)
            self.assertEqual(report['status'], 'success')
            doc = Document(output)
            full_text = '\n'.join(p.text for p in doc.paragraphs)
            self.assertIn('48 mm 法兰盖；承压面', full_text)
            self.assertIn('4 mm 衬里；不承压', full_text)
            self.assertNotIn('====', full_text)
            self.assertNotIn('----', full_text)
            self.assertEqual(report['output_counts']['math'], 2)

    def test_custom_template_preserves_identifiers_and_gets_pagination_guards(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            reference, output = root / 'reference.docx', root / 'output.docx'
            Document().save(reference)
            source = ('# Review\n\n| Parameter | Value |\n|---|---|\n'
                      '| Length | ∠50×50×6，L=2000，两件 |\n\n'
                      '> Formula introduction\n\n$$t=\\frac{PR}{SE-0.6P}$$')
            render_markdown(source, output, reference_doc=reference, pandoc_path=PANDOC)
            doc = Document(output)
            self.assertEqual(doc.tables[0].cell(1, 1).text, '∠50×50×6，L=2000，两件')
            for row in doc.tables[0].rows:
                self.assertIsNotNone(row._tr.find(qn('w:trPr')).find(qn('w:cantSplit')))
            self.assertTrue(all(p.paragraph_format.keep_with_next for c in doc.tables[0].rows[0].cells for p in c.paragraphs))
            intro = next(p for p in doc.paragraphs if p.text == 'Formula introduction')
            self.assertTrue(intro.paragraph_format.keep_with_next)
            self.assertTrue(intro.paragraph_format.keep_together)


if __name__ == '__main__':
    unittest.main()
