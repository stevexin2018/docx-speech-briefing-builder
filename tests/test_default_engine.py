"""Omitting engine selection must actually produce native v2 output."""
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from briefing_v2.engine import BuildError, find_pandoc
from build_briefing import build_docx_worker, build_briefing_concurrent
from render_docx import create_docx_document

try:
    PANDOC = find_pandoc()
except FileNotFoundError:
    PANDOC = None

CORE = Path(__file__).resolve().parents[1]


@unittest.skipUnless(PANDOC, 'Pandoc required for default-engine integration')
class DefaultEngineTests(unittest.TestCase):
    def test_unqualified_python_and_cli_entries_produce_native_math(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / 'source.md'
            text = '# Native default\n\nFormula:\n\n$$t=\\frac{PR}{SE-0.6P}$$'
            source.write_text(text, encoding='utf-8')
            entries = {
                'renderer-api': lambda out: create_docx_document('Default', 'QA', text, out),
                'worker-api': lambda out: build_docx_worker(source, out, 'Default', 'QA'),
                'concurrent-api': lambda out: build_briefing_concurrent(source, out),
            }
            with patch.dict(os.environ, PANDOC_PATH=PANDOC):
                for name, entry in entries.items():
                    output = root / (name + '.docx')
                    entry(output)
                    self.assertEqual(len(Document(output).element.xpath('.//m:f')), 1, name)
                for script, flag in [('render_docx.py', '--output'), ('build_briefing.py', '--docx'),
                                     ('trial_briefing.py', '--docx')]:
                    output = root / (script + '.docx')
                    result = subprocess.run([sys.executable, str(CORE / script), '--input', str(source),
                                             flag, str(output)], capture_output=True, text=True)
                    self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
                    self.assertEqual(len(Document(output).element.xpath('.//m:f')), 1, script)

    def test_missing_pandoc_fails_without_overwriting_or_legacy_fallback(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / 'existing.docx'
            output.write_bytes(b'previous-document')
            with self.assertRaises(BuildError):
                create_docx_document('Default', 'QA', '# Report', output,
                                     pandoc_path=str(Path(directory) / 'missing-pandoc'))
            self.assertEqual(output.read_bytes(), b'previous-document')

    def test_legacy_remains_explicitly_available(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / 'legacy.docx'
            create_docx_document('Compatibility', 'QA', '# Report\n\nBody', output, engine='legacy')
            self.assertIn('Body', '\n'.join(p.text for p in Document(output).paragraphs))
