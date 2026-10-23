import sys
from pathlib import Path
import tempfile
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import docx
from clean_speech_text import clean_speech_text
from render_docx import clean_inline_text, create_docx_document


class QuantityTests(unittest.TestCase):
    def test_quantity_commas(self):
        for source in ('10,mm', '10，mm', '10, mm'):
            self.assertEqual(clean_inline_text(source), '10 mm')
            self.assertEqual(clean_speech_text(source), '10 毫米。')
        self.assertEqual(clean_inline_text('10, apples; 21,120 mm'), '10, apples; 21,120 mm')

    def test_grouped_numbers(self):
        self.assertEqual(clean_speech_text('100,010'), '十万零一十。')
        for source, expected in [('21,120', '两万一千一百二十'), ('21，120', '两万一千一百二十'), ('1,005', '一千零五'), ('200,000', '二十万'), ('100,001', '十万零一'), ('21,120.05', '两万一千一百二十点零五')]:
            self.assertEqual(clean_speech_text(source), expected + '。')
        self.assertIn('两万一千一百二十至两万二千', clean_speech_text('21,120-22,000 mm'))
        self.assertIn('Q W 第 404 点 12 条', clean_speech_text('QW-404.12'))

    def test_powers(self):
        for source in ('mm2', 'mm^2', 'mm^{2}', 'mm²'):
            self.assertEqual(clean_inline_text(source), 'mm²')
            self.assertEqual(clean_speech_text(source), '平方毫米。')
        self.assertEqual(clean_inline_text('10,mm3; cm2; m^3; mm4'), '10 mm³; cm²; m³; mm⁴')
        self.assertEqual(clean_speech_text('N/mm2'), '牛 每 平方毫米。')
        self.assertEqual(clean_speech_text('kg/m3'), '千克 每 立方米。')
        for source, expected in [('cm²', '平方厘米。'), ('km³', '立方千米。'), ('10mm²', '10 平方毫米。'), ('面积mm²', '面积 平方毫米。')]:
            self.assertEqual(clean_speech_text(source), expected)
        for identifier in ('itemm²', 'itemm2', 'mm20', 'K3', 'S2', 'm2', 'unit_mm²', 'mm²_value'):
            self.assertEqual(clean_speech_text(identifier), identifier.replace('_', ' ') + '。')
        self.assertEqual(clean_inline_text('itemm2 mm20 `mm2`'), 'itemm2 mm20 `mm2`')

    def test_imperial_rate_units(self):
        for unit, spoken in (("in/min", "分钟"), ("in./min", "分钟"), ("in/s", "秒"), ("in./s", "秒"), ("in/h", "小时"), ("in./h", "小时")):
            self.assertEqual(clean_speech_text(f"焊接速度 10 {unit}"), f"焊接速度 10 英寸 每 {spoken}。")
        for unit in ("J/in", "J/in."):
            self.assertEqual(clean_speech_text(unit), "焦 每 英寸。")

    def test_document(self):
        with tempfile.TemporaryDirectory() as directory:
            output = str(Path(directory) / 'quantities.docx')
            create_docx_document('计量验证', 'QA', '面积 21,120 mm2\n\n| 项目 | 数值 |\n| --- | --- |\n| 容积 | 10,mm3 |', output)
            document = docx.Document(output)
            self.assertTrue(any('21,120 mm²' in p.text for p in document.paragraphs))
            self.assertTrue(any('10 mm³' in c.text for t in document.tables for r in t.rows for c in r.cells))


if __name__ == '__main__':
    unittest.main()
