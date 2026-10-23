"""Behavioral checks for native math, source protection and failed-build publication."""
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from lxml import etree

from briefing_v2 import BuildError, render_markdown
from briefing_v2.engine import NS, find_pandoc, inspect_docx
from briefing_v2.normalization import normalize_ast
from briefing_v2.source_protection import prepare_source


try:
    PANDOC = find_pandoc()
except FileNotFoundError:
    PANDOC = None


class NormalizationTests(unittest.TestCase):
    def test_path_protection_skips_code_and_mathematics(self):
        source = '`C:\\Temp\\_data`\n\n```text\nC:\\Temp\\_data\n```\n\n    C:\\Temp\\_data\n\n$x_{C:\\Temp}$'
        prepared, diagnostics, count = prepare_source(source, 'report.md')
        self.assertEqual(prepared, source)
        self.assertEqual(count, 0)
        self.assertFalse(diagnostics)

    def test_code_math_and_link_targets_keep_their_content(self):
        ast = {"blocks": [{"t": "Para", "c": [
            {"t": "Code", "c": [["", [], []], r"C:\mm2\报告.md"]},
            {"t": "Math", "c": [{"t": "InlineMath"}, r"x_{mm2}"]},
            {"t": "Link", "c": [["", [], []], [{"t": "Str", "c": "文件"}], ["https://example.test/mm2", ""]]},
        ]}]}
        normalized, analysis, diagnostics = normalize_ast(ast, "source")
        self.assertEqual(normalized, ast)
        self.assertEqual(analysis["source_counts"]["Math"], 1)
        self.assertFalse(diagnostics)

    def test_units_use_native_scripts_without_guessing_model_numbers(self):
        ast = {"blocks": [{"t": "Para", "c": [
            {"t": "Str", "c": "mm2"}, {"t": "Space"}, {"t": "Str", "c": "m2"},
            {"t": "Space"}, {"t": "Str", "c": "K3"}, {"t": "Space"}, {"t": "Str", "c": "mm20"},
        ]}]}
        normalized, analysis, _ = normalize_ast(ast, "mm2 m2 K3 mm20")
        inline = normalized["blocks"][0]["c"]
        self.assertIn({"t": "Superscript", "c": [{"t": "Str", "c": "2"}]}, inline)
        for identifier in ("m2", "K3", "mm20"):
            self.assertIn({"t": "Str", "c": identifier}, inline)
        self.assertEqual(ast["blocks"][0]["c"][0]["c"], "mm2")

    def test_ambiguous_source_position_is_not_invented(self):
        ast = {"blocks": [{"t": "Para", "c": [{"t": "Str", "c": "5^"}]}]}
        _, _, diagnostics = normalize_ast(ast, "5^ and 5^", "report.md")
        self.assertEqual(diagnostics[0].code, "AMBIGUOUS_POWER")
        self.assertIsNone(diagnostics[0].line)
        self.assertIsNotNone(diagnostics[0].node)
        self.assertEqual(diagnostics[0].excerpt, "5^")


class MissingDependencyTests(unittest.TestCase):
    def test_failure_has_structured_diagnostic_and_no_docx(self):
        with tempfile.TemporaryDirectory() as folder:
            output, diagnostic = Path(folder) / "out.docx", Path(folder) / "diagnostic.json"
            with patch("briefing_v2.engine.find_pandoc", side_effect=FileNotFoundError("missing Pandoc")):
                with self.assertRaises(BuildError) as raised:
                    render_markdown("# Report", output, diagnostics_path=diagnostic)
            self.assertFalse(output.exists())
            self.assertEqual(raised.exception.report["status"], "failed")
            self.assertEqual(json.loads(diagnostic.read_text(encoding="utf-8"))["diagnostics"][0]["code"], "DEPENDENCY_NOT_FOUND")


@unittest.skipUnless(PANDOC, "Set PANDOC_PATH to run Pandoc integration tests")
class PandocIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.root = Path(self.folder.name)
        self.output = self.root / "报告.docx"

    def tearDown(self):
        self.folder.cleanup()

    def render(self, text, **kwargs):
        return render_markdown(text, self.output, pandoc_path=PANDOC, **kwargs)

    def test_complex_formulas_and_table_formula_are_native(self):
        expressions = [r"\sqrt[3]{x}", r"\begin{bmatrix}a&b\\c&d\end{bmatrix}",
                       r"f(x)=\begin{cases}x&x\ge0\\-x&x<0\end{cases}",
                       r"\begin{aligned}a&=b\\c&=d\end{aligned}", r"a < b", r"\text{A \& B}"]
        source = "# 测试报告\n\n" + "\n\n".join("$$\n" + expression + "\n$$" for expression in expressions)
        source += "\n\n## 参数表\n\n| 参数 | 公式 |\n| --- | --- |\n| 厚度 | $\\frac{P}{S}$ |\n"
        report = self.render(source)
        self.assertEqual(report["output_counts"]["math"], 7)
        self.assertEqual(report["output_counts"]["table_math"], 1)
        self.assertEqual(report["output_counts"]["headings"], 2)
        self.assertEqual(report["output_counts"]["repeat_headers"], 1)
        with ZipFile(self.output) as archive:
            root = etree.fromstring(archive.read("word/document.xml"))
        self.assertEqual(len(root.xpath('//m:rad', namespaces=NS)), 1)
        self.assertGreaterEqual(len(root.xpath('//m:m | //m:eqArr', namespaces=NS)), 3)

    def test_unknown_macro_is_located_and_does_not_replace_existing_output(self):
        self.output.write_bytes(b"previous valid output sentinel")
        diagnostic = self.root / "diagnostics.json"
        with self.assertRaises(BuildError) as raised:
            self.render("# Report\n\n$$\\DefinitelyUnsupportedMacro{x}$$", diagnostics_path=diagnostic, source_name="report.md")
        self.assertEqual(self.output.read_bytes(), b"previous valid output sentinel")
        errors = raised.exception.diagnostics
        failure = next(item for item in errors if item["code"] == "MATH_CONVERSION_FAILED")
        self.assertEqual(failure["line"], 3)
        self.assertTrue(failure["node"])
        self.assertEqual(json.loads(diagnostic.read_text(encoding="utf-8"))["status"], "failed")

    def test_code_and_paths_are_exact_and_units_get_native_superscripts(self):
        path = r"C:\Program Files\mm2\_temp\报告.md"
        source = "# Report\n\n10 mm2、10 m3；m2 K3 mm20。\n\n`" + path + "`\n\n```text\necho " + path + "\nnext_line\n```\n"
        report = self.render(source)
        self.assertEqual(report["protected_items"], 2)
        text = inspect_docx(self.output)["text"]
        self.assertIn(path, text)
        self.assertIn("m2 K3 mm20", text)
        with ZipFile(self.output) as archive:
            root = etree.fromstring(archive.read("word/document.xml"))
        self.assertTrue(root.xpath('//w:vertAlign[@w:val="superscript"]', namespaces=NS))
        self.assertTrue(root.xpath('//w:r[w:rPr/w:vertAlign[@w:val="superscript"]]/w:t[text()="3"]', namespaces=NS))

    def test_both_backslash_math_delimiters_and_ambiguous_temperature(self):
        report = self.render("# Report\n\n\\(x_i^2\\)\n\n\\[\\frac{a}{b}\\]\n\n条件为 5^。")
        self.assertEqual(report["output_counts"]["math"], 2)
        self.assertTrue(any(item["code"] == "AMBIGUOUS_POWER" for item in report["diagnostics"]))
        self.assertNotIn("°C", inspect_docx(self.output)["text"])

    def test_raw_latex_is_reported_instead_of_disappearing(self):
        with self.assertRaises(BuildError) as raised:
            self.render(r"\frac{a}{b}")
        self.assertTrue(any(item["code"] == "UNDELIMITED_LATEX" for item in raised.exception.diagnostics))

    def test_bare_symbols_are_kept_without_guessing_a_complete_equation(self):
        report = self.render(r"条件为 (0<\Delta t/\tau<2)。")
        text = inspect_docx(self.output)["text"]
        self.assertIn("Δ", text)
        self.assertIn("τ", text)
        self.assertEqual(report["output_counts"]["math"], 0)
        self.assertTrue(any(item["code"] == "BARE_SYMBOL_NORMALIZED" for item in report["diagnostics"]))

    def test_bare_windows_path_cannot_be_read_as_latex_or_units(self):
        literal = r"C:\Temp\_data\mm2.md"
        report = self.render("文件位于 " + literal + "。")
        self.assertIn(literal, inspect_docx(self.output)["text"])
        self.assertEqual(report["protected_bare_paths"], 1)

    def test_quoted_path_with_spaces_is_preserved(self):
        literal = r"C:\Program Files\_data\mm2.md"
        self.render('文件位于 "' + literal + '"。')
        self.assertIn(literal, inspect_docx(self.output)["text"])

    def test_unclosed_display_or_parenthesis_math_is_reported(self):
        for source in (r"\(x_i", "$$x_i", r"\[x_i"):
            with self.subTest(source=source), self.assertRaises(BuildError) as raised:
                self.render(source)
            self.assertTrue(any(item["code"] == "UNCLOSED_MATH" for item in raised.exception.diagnostics))

    def test_reference_commands_are_not_silently_discarded(self):
        with self.assertRaises(BuildError) as raised:
            self.render(r"$$x=1\label{eq:one}$$")
        self.assertTrue(any(item["code"] == "UNSUPPORTED_LATEX_COMMAND" for item in raised.exception.diagnostics))

    def test_br_in_table_is_a_native_break(self):
        self.render("| 项目 | 内容 |\n| --- | --- |\n| A | 第一行<br>第二行 |\n")
        text = inspect_docx(self.output)["text"]
        self.assertIn("第一行\n第二行", text)

    def test_protection_guard_catches_changed_content(self):
        from briefing_v2.template import enhance_document
        def corrupt(path, *args, **kwargs):
            enhance_document(path, *args, **kwargs)
            document = Document(path)
            for paragraph in document.paragraphs:
                if "KEEP_THIS_CODE" in paragraph.text:
                    paragraph.text = "changed"
            document.save(path)
        with patch("briefing_v2.engine.enhance_document", side_effect=corrupt):
            with self.assertRaises(BuildError) as raised:
                self.render("`KEEP_THIS_CODE`")
        self.assertFalse(self.output.exists())
        self.assertTrue(any(item["code"] == "PROTECTED_CONTENT_CHANGED" for item in raised.exception.diagnostics))

    def test_diagnostics_cannot_overwrite_reference(self):
        template = self.root / "template.docx"
        Document().save(template)
        before = template.read_bytes()
        with self.assertRaises(BuildError):
            self.render("# Report", reference_doc=template, diagnostics_path=template)
        self.assertEqual(template.read_bytes(), before)

    def test_template_header_math_does_not_count_as_a_missing_source_formula(self):
        from docx.oxml import parse_xml
        template = self.root / "template.docx"
        document = Document()
        document.sections[0].header.paragraphs[0]._p.append(parse_xml(
            '<m:oMath xmlns:m="' + NS['m'] + '"><m:r><m:t>x</m:t></m:r></m:oMath>'))
        document.save(template)
        before = template.read_bytes()
        report = self.render("# Report\n\n正文。", reference_doc=template)
        self.assertEqual(report['output_counts']['math'], 0)
        self.assertEqual(template.read_bytes(), before)
        with ZipFile(self.output) as archive:
            self.assertIn(b'oMath', archive.read('word/header1.xml'))

    def test_title_option_adds_a_title_only_when_source_has_none(self):
        self.render("正文。", title="报告标题")
        document = Document(self.output)
        self.assertEqual(document.paragraphs[0].text, "报告标题")
        self.assertEqual(document.paragraphs[0].style.name, "Title")


if __name__ == "__main__":
    unittest.main()
